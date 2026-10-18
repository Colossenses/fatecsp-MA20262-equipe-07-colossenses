from openpyxl import load_workbook
from numbers import Number

arquivo = "transport.xlsx"
arquivo_saida = "transport_atualizado.xlsx"

wb = load_workbook(arquivo)
ws = wb["Trip"]

# Localizar a coluna service_id
coluna = None

for celula in ws[1]:
    if celula.value == "service_id":
        coluna = celula.column
        break

if coluna is None:
    raise ValueError("Coluna service_id não encontrada.")

for linha in range(2, ws.max_row + 1):
    celula = ws.cell(row=linha, column=coluna)
    valor = celula.value

    if isinstance(valor, str) and not isinstance(valor, bool):
        celula.value = (
            f'=_xlfn.XLOOKUP("{valor}",Calendar[service_id],Calendar[service_id],"Não encontrado")'
        )

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"

wb.save(arquivo_saida)

print("Arquivo salvo:", arquivo_saida)